"""Sovereign AI Workbench -- API surface.

Everything the UI needs, and nothing it does not: route a request, stream the
agent trace, prove containment.
"""
from __future__ import annotations

import asyncio
import json
import os
from contextlib import asynccontextmanager
from pathlib import Path

from fastapi import Depends, FastAPI, Form, File, HTTPException, Header, UploadFile
from fastapi.responses import StreamingResponse, FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles

from . import audit
from .agent import Agent
from .auth import ROLES, USERS, User
from .sessions import CONVERSATIONS, Turn
from .egress import EgressMonitor
from .kb import KB
from .llm import LLM, load_dotenv
from .router import Router

ROOT = Path(__file__).resolve().parent.parent
UPLOADS = ROOT / "data" / "uploads"
OUT = ROOT / "data" / "out"
MODE = os.environ.get("MODE", "prototype")

load_dotenv(ROOT / ".env")
router = Router(mode=MODE)
monitor = EgressMonitor(allowlist=router.mode_cfg.get("egress_allow") or [], mode=MODE)


@asynccontextmanager
async def lifespan(app: FastAPI):
    await monitor.start()
    # A fresh clone has documents but no index. Build it rather than serving an
    # empty knowledge base, which would make every task ungrounded on first run.
    if not KB.load() and any((ROOT / "data" / "kb").glob("*.md")):
        stats = await asyncio.to_thread(KB.build)
        audit.record("kb.built_on_startup", **stats)
    yield


app = FastAPI(title="Sovereign AI Workbench", lifespan=lifespan)


@app.middleware("http")
async def no_store_api(request, call_next):
    """Nothing under /api may be cached by the browser.

    API responses carried no Cache-Control and varied only on Origin, so a
    browser could answer a later request for the same URL from its cache --
    including a request with no token, after sign-out. On a workstation shared
    between shifts that replays one user's documents to whoever sits down next.
    The interface's own assets are content-hashed and stay cacheable.
    """
    response = await call_next(request)
    if request.url.path.startswith("/api/"):
        response.headers["Cache-Control"] = "no-store"
        response.headers["Pragma"] = "no-cache"
    return response


# ---------------------------------------------------------------------------
# Identity. Every action carries a name so the audit log is worth keeping.
# ---------------------------------------------------------------------------
def current_user(authorization: str = Header(default="")) -> User:
    token = authorization.removeprefix("Bearer ").strip()
    user = USERS.whoami(token)
    if not user:
        raise HTTPException(status_code=401, detail="sign in required")
    return user


def needs(capability: str):
    def dep(user: User = Depends(current_user)) -> User:
        if not user.can(capability):
            raise HTTPException(
                status_code=403,
                detail=f"your role ({user.role}) cannot {capability}")
        return user
    return dep


@app.post("/api/login")
async def login(username: str = Form(...), password: str = Form(...)):
    token = USERS.login(username, password)
    if not token:
        # One message for both cases: saying which was wrong tells an attacker
        # which usernames exist.
        audit.record("login.failed", username=username[:64])
        raise HTTPException(status_code=401, detail="wrong username or password")
    user = USERS.whoami(token)
    audit.record("login", user=user.username, role=user.role)
    return {"token": token, "user": user.public()}


@app.post("/api/logout")
async def logout(authorization: str = Header(default="")):
    USERS.logout(authorization.removeprefix("Bearer ").strip())
    return {"ok": True}


@app.get("/api/me")
async def me(user: User = Depends(current_user)):
    return {"user": user.public(), "roles": {r: sorted(c) for r, c in ROLES.items()}}


@app.post("/api/me/password")
async def change_password(current: str = Form(...), new: str = Form(...),
                          user: User = Depends(current_user)):
    from .auth import verify_password
    if not verify_password(current, user.password):
        raise HTTPException(status_code=403, detail="current password is wrong")
    try:
        USERS.set_password(user.username, new)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    audit.record("password.changed", user=user.username)
    return {"ok": True}


# ---------------------------------------------------------------------------
# Administration
# ---------------------------------------------------------------------------
@app.get("/api/users")
async def list_users(user: User = Depends(needs("manage_users"))):
    return {"users": [u.public() for u in USERS.users.values()],
            "roles": {r: sorted(c) for r, c in ROLES.items()}}


@app.post("/api/users")
async def add_user(username: str = Form(...), password: str = Form(...),
                   role: str = Form("engineer"), display: str = Form(""),
                   user: User = Depends(needs("manage_users"))):
    try:
        created = USERS.add(username, password, role, display)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    audit.record("user.created", by=user.username, username=created.username,
                 role=created.role)
    return {"user": created.public()}


@app.post("/api/users/{username}/role")
async def set_role(username: str, role: str = Form(...),
                   user: User = Depends(needs("manage_users"))):
    if username not in USERS.users:
        raise HTTPException(status_code=404, detail="no such user")
    try:
        USERS.set_role(username, role)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    audit.record("user.role_changed", by=user.username, username=username, role=role)
    return {"user": USERS.users[username].public()}


@app.delete("/api/users/{username}")
async def remove_user(username: str, user: User = Depends(needs("manage_users"))):
    try:
        removed = USERS.remove(username)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    if not removed:
        raise HTTPException(status_code=404, detail="no such user")
    audit.record("user.removed", by=user.username, username=username)
    return {"ok": True}


@app.get("/api/status")
async def status():
    return {
        "mode": MODE,
        "endpoint": router.mode_cfg["endpoint"],
        "egress": monitor.snapshot(),
        "models": [
            {"id": m["id"], "tier": m["tier"], "caps": m["caps"],
             "name": router.resolve(m), "why": m["why"],
             "max_tokens": m.get("max_tokens")}
            for m in router.models
        ],
        "rules": [{"name": r["name"], "why": r["why"], "if": r.get("if", {}),
                   "then": r.get("then", {})} for r in router.routing_rules],
    }


@app.post("/api/registry/reload")
async def reload_registry(user: User = Depends(needs("manage_models"))):
    """Re-read models.yaml. This is the 'add a model without redesign' proof."""
    before = {m["id"] for m in router.models}
    router.reload()
    after = {m["id"] for m in router.models}
    return {"added": sorted(after - before), "removed": sorted(before - after),
            "models": sorted(after)}


@app.post("/api/route")
async def preview_route(prompt: str = Form(...), has_image: bool = Form(False),
                        user: User = Depends(current_user)):
    """Routing decision without executing it -- used by the UI's router panel."""
    return router.route(prompt, has_image=has_image).as_dict()


@app.post("/api/tripwire")
async def tripwire(target: str = Form("https://api.openai.com/v1/models"),
                   user: User = Depends(current_user)):
    r = await monitor.tripwire(target)
    audit.record("tripwire", user=user.username, **r)
    return r


@app.get("/api/egress/stream")
async def egress_stream(user: User = Depends(current_user)):
    """Live tail of observed outbound traffic.

    Authenticated by header like everything else. EventSource cannot send
    headers, and a token in the query string would be written to the access
    log, so the interface reads this with fetch() instead.
    """
    q: asyncio.Queue = asyncio.Queue(maxsize=100)
    monitor.subscribers.add(q)

    async def gen():
        try:
            yield f"data: {json.dumps({'type': 'snapshot', **monitor.snapshot()})}\n\n"
            while True:
                try:
                    ev = await asyncio.wait_for(q.get(), timeout=15)
                    yield f"data: {json.dumps({'type': 'egress', **ev.as_dict()})}\n\n"
                except asyncio.TimeoutError:
                    yield ": keepalive\n\n"
        finally:
            monitor.subscribers.discard(q)

    return StreamingResponse(gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache",
                                      "X-Accel-Buffering": "no"})


@app.post("/api/run")
async def run(prompt: str = Form(...), has_image: bool = Form(False),
              attachment: str = Form(""), conversation: str = Form(""),
              user: User = Depends(needs("run"))):
    """Execute a task, streaming every step as it happens.

    The trace is the product as much as the answer is: an engineer approving a
    note needs to see which SOP clause was retrieved and what was computed, not
    just the conclusion.
    """
    q: asyncio.Queue = asyncio.Queue()
    # An id belonging to someone else resolves to a fresh conversation rather
    # than an error, so the endpoint never confirms that the id exists.
    conv = CONVERSATIONS.get_or_create(user.username, conversation or None)
    # A follow-up about an attached report should not need it re-attached.
    carried = attachment or conv.carried_attachment() or ""
    audit.record("task.received", user=user.username, prompt=prompt[:500],
                 mode=MODE, attachment=carried or None, conversation=conv.id,
                 turn=len(conv.turns) + 1)

    # The attachment is passed through separately rather than pasted into the
    # prompt. Appending guidance text to the prompt fed words like "drawing"
    # into the router, which then sent a scanned report to the drawing model.

    async def gen():
        # Steps are emitted from the event-loop thread, so queue them directly.
        # Going through call_soon_threadsafe would schedule them *after* the
        # sentinel that drive() puts in synchronously, and a fast run would
        # stream its result with no visible trace at all.
        agent = Agent(router, LLM(router), on_step=q.put_nowait)

        async def drive():
            try:
                return await agent.run(prompt, attachment=carried or None,
                                       history=conv.history())
            finally:
                q.put_nowait(None)

        task = asyncio.create_task(drive())
        steps: list[dict] = []
        yield f"data: {json.dumps({'type': 'conversation', 'id': conv.id})}\n\n"
        while True:
            item = await q.get()
            if item is None:
                break
            steps.append(item.as_dict())
            d = item.as_dict()
            audit.record(f"step.{d.pop('kind')}", **d)
            yield f"data: {json.dumps({'type': 'step', **item.as_dict()}, default=str)}\n\n"
        try:
            final = await task
        except Exception as e:
            audit.record("task.failed", error=f"{type(e).__name__}: {e}")
            yield f"data: {json.dumps({'type': 'error', 'error': f'{type(e).__name__}: {e}'})}\n\n"
            return
        conv_after = CONVERSATIONS.add_turn(user.username, conv.id, Turn(
            prompt=prompt, answer=final.get("answer", ""),
            evidence=final.get("evidence", []),
            deliverables=final.get("deliverables", []),
            attachment=carried or None, decision=final.get("decision", {}),
            steps=steps, verdict=final.get("verdict")))
        audit.record("task.completed", user=user.username,
                     verdict=final.get("verdict"),
                     model=final["decision"].get("model_name"), conversation=conv.id,
                     deliverables=[d["path"] for d in final["deliverables"]])
        yield f"data: {json.dumps({'type': 'final', 'conversation': conv.id, 'title': conv_after.title, **final}, default=str)}\n\n"

    return StreamingResponse(gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache",
                                      "X-Accel-Buffering": "no"})


@app.post("/api/kb/build")
async def kb_build(user: User = Depends(needs("manage_kb"))):
    stats = await asyncio.to_thread(KB.build)
    audit.record("kb.rebuilt", user=user.username, **stats)
    return stats


@app.get("/api/kb/search")
async def kb_search(q: str, k: int = 6, user: User = Depends(needs("read_kb"))):
    if KB.vectors is None:
        KB.load()
    return {"query": q, "passages": KB.search(q, k=k)}


@app.get("/api/kb/documents")
async def kb_documents(user: User = Depends(needs("read_kb"))):
    if KB.vectors is None:
        KB.load()
    docs: dict[str, int] = {}
    for c in KB.chunks:
        docs[c.doc] = docs.get(c.doc, 0) + 1
    return {"chunks": len(KB.chunks),
            "documents": [{"name": d, "chunks": n} for d, n in sorted(docs.items())]}


# ---------------------------------------------------------------------------
# Conversations. Every call is scoped to the signed-in user by the store itself;
# another user's conversation is a 404, never a 403, so an id is never
# confirmed to exist.
# ---------------------------------------------------------------------------
@app.get("/api/conversations")
async def list_conversations(user: User = Depends(current_user)):
    return {"conversations": CONVERSATIONS.list(user.username)}


@app.get("/api/conversations/{cid}")
async def get_conversation(cid: str, user: User = Depends(current_user)):
    c = CONVERSATIONS.get(user.username, cid)
    if not c:
        raise HTTPException(status_code=404, detail="no such conversation")
    return c.full()


@app.patch("/api/conversations/{cid}")
async def rename_conversation(cid: str, title: str = Form(...),
                              user: User = Depends(current_user)):
    c = CONVERSATIONS.rename(user.username, cid, title)
    if not c:
        raise HTTPException(status_code=404, detail="no such conversation")
    return c.summary()


@app.delete("/api/conversations/{cid}")
async def delete_conversation(cid: str, user: User = Depends(current_user)):
    if not CONVERSATIONS.delete(user.username, cid):
        raise HTTPException(status_code=404, detail="no such conversation")
    audit.record("conversation.deleted", user=user.username, conversation=cid)
    return {"ok": True}


@app.get("/api/audit")
async def audit_log(n: int = 150, user: User = Depends(needs("read_audit"))):
    return {"entries": audit.tail(n)}


@app.post("/api/upload")
async def upload(file: UploadFile = File(...),
                 user: User = Depends(needs("upload"))):
    UPLOADS.mkdir(parents=True, exist_ok=True)
    name = Path(file.filename or "upload.bin").name
    if not name or name.startswith("."):
        return JSONResponse({"error": "unusable filename"}, status_code=400)
    dest = UPLOADS / name
    data = await file.read()
    # 32MB: a scanned report is a few MB, and an unbounded upload into a
    # container with a small disk is a denial of service with extra steps.
    if len(data) > 32 * 1024 * 1024:
        return JSONResponse({"error": "file larger than 32MB"}, status_code=413)
    UPLOADS.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(data)
    audit.record("upload", user=user.username, name=dest.name,
                 bytes=dest.stat().st_size)
    return {"path": str(dest.relative_to(ROOT)), "bytes": dest.stat().st_size}


@app.get("/api/download/{name}")
async def download(name: str, user: User = Depends(current_user)):
    p = OUT / Path(name).name
    if not p.exists():
        return JSONResponse({"error": "not found"}, status_code=404)
    return FileResponse(p, filename=p.name)


@app.get("/api/files/{name}/outline")
async def file_outline(name: str, user: User = Depends(current_user)):
    """What a deliverable contains, for the file card: its headings, the
    citations it carries, and its size -- read from the file itself, the same
    way verification reads it back."""
    from .agent import cited_ids
    p = OUT / Path(name).name
    if not p.exists():
        raise HTTPException(status_code=404, detail="no such file")
    out: dict = {"name": p.name, "bytes": p.stat().st_size,
                 "kind": p.suffix.lstrip(".").lower()}
    if p.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(p)
        text = "\n".join([x.text for x in doc.paragraphs] +
                         [c.text for t in doc.tables for r in t.rows for c in r.cells])
        out |= {"title": next((x.text for x in doc.paragraphs
                               if x.style.name == "Title" and x.text.strip()), p.stem),
                "headings": [x.text for x in doc.paragraphs
                             if x.style.name.startswith("Heading") and x.text.strip()],
                "tables": len(doc.tables),
                "citations": sorted(cited_ids(text))}
    elif p.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True)
        out |= {"sheets": [{"name": ws.title, "rows": ws.max_row,
                            "columns": ws.max_column} for ws in wb.worksheets]}
    return out


class SPA(StaticFiles):
    """Serve the built interface, falling back to index.html.

    The interface routes on the client -- /c/<conversation> is a real address a
    user can bookmark or reload -- so an unknown path that is not an API call
    must return the app rather than a 404.
    """
    async def get_response(self, path: str, scope):
        # Starlette's exception, not FastAPI's: StaticFiles raises the base
        # class, and FastAPI's HTTPException is a subclass of it, so catching
        # FastAPI's would let every 404 through and the fallback never fire.
        from starlette.exceptions import HTTPException as StarletteHTTPException
        try:
            return await super().get_response(path, scope)
        except StarletteHTTPException as e:
            if e.status_code == 404 and not path.startswith("api/"):
                return await super().get_response("index.html", scope)
            raise


app.mount("/", SPA(directory=ROOT / "static", html=True), name="static")
