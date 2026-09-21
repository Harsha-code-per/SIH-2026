"""Deliverables must be documents a section head can read, not markdown dumps."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document
from app.tools import write_docx, _split_markdown_tables, ROOT

MD_BODY = """Key measurements:

| Point | Horiz | Vert |
|-------|-------|------|
| Drive end | 8.2 | 7.4 |
| Non-drive end | 6.9 | 6.1 |

Temperature is 79 C against a 61 C baseline."""


def test_markdown_tables_are_split_from_prose():
    parts = _split_markdown_tables(MD_BODY)
    kinds = [k for k, _ in parts]
    assert kinds == ["text", "table", "text"], kinds
    rows = parts[1][1]
    assert rows[0] == ["Point", "Horiz", "Vert"]
    assert rows[1] == ["Drive end", "8.2", "7.4"]
    assert len(rows) == 3, "the |---| separator must not become a row"


def test_written_document_has_a_real_table():
    r = write_docx("Approval Note",
                   [{"heading": "Observation", "body": MD_BODY,
                     "citations": ["Maintenance_SOP_v7.md#7"]}],
                   filename="_test_note.docx")
    doc = Document(ROOT / r["path"])
    assert len(doc.tables) == 1, "markdown table should become a Word table"
    t = doc.tables[0]
    assert t.cell(0, 0).text == "Point"
    assert t.cell(1, 1).text == "8.2"
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "|---" not in text and "| Point |" not in text, "raw markdown leaked into prose"
    assert "Maintenance_SOP_v7.md#7" in text, "citation missing"
    (ROOT / r["path"]).unlink()


def test_ragged_rows_do_not_crash():
    r = write_docx("T", [{"body": "| a | b | c |\n|---|---|---|\n| 1 |\n| 1 | 2 | 3 |"}],
                   filename="_test_ragged.docx")
    doc = Document(ROOT / r["path"])
    assert doc.tables[0].cell(1, 2).text == ""
    (ROOT / r["path"]).unlink()


def test_prose_without_tables_is_untouched():
    parts = _split_markdown_tables("Just a sentence.\n\nAnd another.")
    assert parts == [("text", "Just a sentence.\n\nAnd another.")]


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_"):
            fn()
            print(f"  ok  {name}")
    print("\ndeliverables: all checks passed")
